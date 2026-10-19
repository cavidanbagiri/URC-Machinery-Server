# utils/seed_maintenance.py
"""Seed periodic maintenance plans from the Excel workbook.

Usage (from the ``server`` directory, with the virtualenv active)::

    python -m utils.seed_maintenance --dry-run
    python -m utils.seed_maintenance
    python -m utils.seed_maintenance --overwrite

The workbook is read from ``static_datas/peryodik_bakim.xlsx`` by default.
Plans are matched to machines by ``all_machines.identification_no``; rows whose
machine is not present in the database are reported and skipped. Existing plans
are left untouched unless ``--overwrite`` is given.

The ``Liste`` sheet defines the usage unit per machine (``Одометр Км`` -> km,
``Моточас`` -> hour). The ``Periyodik Bakım`` sheet provides the interval, the
last completed maintenance counter/date, and the current counter. A ``Pert``
value in the last-maintenance-date column marks the plan as inactive.
"""

from __future__ import annotations

import argparse
import asyncio
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Optional

import openpyxl
from sqlalchemy import select

from database.setup import SessionLocal
from models.machines_model import AllMachineModel
from models.maintenance_model import MachineMaintenanceSettingModel
from models.user_model import UserModel  # noqa: F401  (mapper dependency)

SERVER_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = SERVER_DIR / "static_datas" / "peryodik_bakim.xlsx"

LIST_SHEET = "Liste"
PLAN_SHEET = "Periyodik Bakım"
PLAN_FIRST_DATA_ROW = 4

# Column indexes (one based) in the "Periyodik Bakım" sheet.
COL_ID = 4
COL_LAST_COUNTER = 7
COL_LAST_DATE = 8
COL_CURRENT_COUNTER = 9
COL_INTERVAL = 11

EXCEL_EPOCH = date(1899, 12, 30)

KM_PREFIXES = ("BA", "BUS", "MB", "ST", "TR", "VD", "VH")
HOUR_PREFIXES = ("BL", "EBP", "EXL", "FL", "FLM", "ML", "TK", "VM", "COM", "IJ", "JD")


def _normalize_id(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text.upper() if text else None


def _parse_number(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace("\u00a0", "")
    if not text or text.casefold() in {"-", "pert", "n/a"}:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _parse_date(value: Any) -> tuple[Optional[date], bool]:
    """Return ``(date, is_pert)`` for a last-maintenance cell."""

    if value is None:
        return None, False
    if isinstance(value, datetime):
        return value.date(), False
    if isinstance(value, date):
        return value, False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return EXCEL_EPOCH + timedelta(days=int(value)), False

    text = str(value).strip()
    if not text:
        return None, False
    if "pert" in text.casefold():
        return None, True

    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date(), False
        except ValueError:
            continue
    print(f"  ! Unrecognised date value ignored: {value!r}")
    return None, False


def _usage_unit_from_label(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().casefold()
    if "одометр" in text or text == "km" or "км" in text:
        return "km"
    if "моточас" in text or "мото-час" in text or text == "hour" or "час" in text:
        return "hour"
    return None


def _infer_usage_unit(machine_id: str) -> str:
    prefix = machine_id.split("-", 1)[0].upper()
    if prefix in KM_PREFIXES:
        return "km"
    if prefix in HOUR_PREFIXES:
        return "hour"
    return "km"


def _cell(ws, row: int, column: int) -> Any:
    return ws.cell(row=row, column=column).value


def read_usage_units(workbook) -> dict[str, str]:
    if LIST_SHEET not in workbook.sheetnames:
        raise SystemExit(f"Sheet not found: {LIST_SHEET}")

    ws = workbook[LIST_SHEET]
    units: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        machine_id = _normalize_id(_cell(ws, row, 3))
        if not machine_id:
            continue
        unit = _usage_unit_from_label(_cell(ws, row, 5))
        if unit:
            units[machine_id] = unit
    return units


def read_plans(workbook, units: dict[str, str]) -> list[dict[str, Any]]:
    if PLAN_SHEET not in workbook.sheetnames:
        raise SystemExit(f"Sheet not found: {PLAN_SHEET}")

    ws = workbook[PLAN_SHEET]
    plans: list[dict[str, Any]] = []

    for row in range(PLAN_FIRST_DATA_ROW, ws.max_row + 1):
        machine_id = _normalize_id(_cell(ws, row, COL_ID))
        if not machine_id:
            continue

        interval = _parse_number(_cell(ws, row, COL_INTERVAL))
        if interval is None or interval <= 0:
            print(f"  ! {machine_id} (row {row}): invalid interval, skipped")
            continue

        last_date, is_pert = _parse_date(_cell(ws, row, COL_LAST_DATE))
        usage_unit = units.get(machine_id)
        if usage_unit is None:
            usage_unit = _infer_usage_unit(machine_id)
            print(
                f"  ! {machine_id}: usage unit missing in '{LIST_SHEET}', "
                f"inferred as '{usage_unit}'"
            )

        plans.append(
            {
                "machine_id": machine_id,
                "row": row,
                "usage_unit": usage_unit,
                "interval_value": interval,
                "last_counter": _parse_number(_cell(ws, row, COL_LAST_COUNTER)),
                "last_maintenance_date": last_date,
                "current_counter": _parse_number(
                    _cell(ws, row, COL_CURRENT_COUNTER)
                ),
                "is_active": not is_pert,
            }
        )

    return plans


def _print_summary(
    plans: Iterable[dict[str, Any]],
    *,
    missing_machines: list[str],
    existing_machines: list[str],
) -> None:
    plans = list(plans)
    active = sum(1 for plan in plans if plan["is_active"])
    print()
    print(f"Plans found in Excel : {len(plans)}")
    print(f"  active             : {active}")
    print(f"  inactive (Pert)    : {len(plans) - active}")
    print(f"Matched in database  : {len(plans) - len(missing_machines)}")
    print(f"Missing in database  : {len(missing_machines)}")
    if missing_machines:
        print("  " + ", ".join(missing_machines))
    if existing_machines:
        print(f"Already have a plan  : {len(existing_machines)}")
        print("  " + ", ".join(existing_machines))


async def seed_database(
    plans: list[dict[str, Any]],
    *,
    dry_run: bool,
    overwrite: bool,
) -> None:
    async with SessionLocal() as session:
        machine_rows = (
            await session.execute(
                select(AllMachineModel.id, AllMachineModel.identification_no)
            )
        ).all()
        machine_ids = {
            normalized: machine_id
            for machine_id, identification_no in machine_rows
            if (normalized := _normalize_id(identification_no))
        }

        setting_rows = (
            await session.execute(
                select(
                    MachineMaintenanceSettingModel.id,
                    MachineMaintenanceSettingModel.machine_id,
                )
            )
        ).all()
        settings_by_machine = {
            machine_id: setting_id for setting_id, machine_id in setting_rows
        }

        missing_machines: list[str] = []
        existing_machines: list[str] = []
        created = 0
        updated = 0
        skipped = 0

        for plan in plans:
            machine_id = machine_ids.get(plan["machine_id"])
            if machine_id is None:
                missing_machines.append(plan["machine_id"])
                continue

            setting_id = settings_by_machine.get(machine_id)
            if setting_id is not None and not overwrite:
                existing_machines.append(plan["machine_id"])
                skipped += 1
                continue

            if setting_id is None:
                setting = MachineMaintenanceSettingModel(machine_id=machine_id)
                session.add(setting)
                created += 1
            else:
                setting = await session.get(
                    MachineMaintenanceSettingModel, setting_id
                )
                updated += 1

            setting.usage_unit = plan["usage_unit"]
            setting.interval_value = plan["interval_value"]
            setting.last_counter = plan["last_counter"]
            setting.last_maintenance_date = plan["last_maintenance_date"]
            setting.current_counter = plan["current_counter"]
            # The workbook has no separate date for the current counter; keep
            # the value without inventing a date.
            setting.current_counter_date = None
            setting.is_active = plan["is_active"]

        _print_summary(
            plans,
            missing_machines=missing_machines,
            existing_machines=existing_machines,
        )
        print(f"Created              : {created}")
        print(f"Updated              : {updated}")
        print(f"Skipped (existing)   : {skipped}")

        if dry_run:
            await session.rollback()
            print("\nDry run: no changes were saved.")
            return

        await session.commit()
        print("\nMaintenance plans saved.")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Seed periodic maintenance plans from the Excel workbook."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help=f"Path to the .xlsx file (default: {DEFAULT_WORKBOOK})",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Read the workbook and report what would change without saving.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Overwrite plans that already exist in the database.",
    )
    args = parser.parse_args()

    workbook_path: Path = args.workbook
    if not workbook_path.is_absolute():
        workbook_path = SERVER_DIR / workbook_path
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    print(f"Reading workbook: {workbook_path}")
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    units = read_usage_units(workbook)
    plans = read_plans(workbook, units)

    try:
        asyncio.run(
            seed_database(plans, dry_run=args.dry_run, overwrite=args.overwrite)
        )
    except KeyboardInterrupt:
        sys.exit(130)


if __name__ == "__main__":
    main()


