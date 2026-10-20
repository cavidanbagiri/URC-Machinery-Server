# utils/seed_machines.py
"""Seed the machine/equipment fleet from the Excel workbook.

Usage (from the ``server`` directory, with the virtualenv active)::

    python -m utils.seed_machines --dry-run
    python -m utils.seed_machines
    python -m utils.seed_machines --create-lookups
    python -m utils.seed_machines --overwrite

The workbook is read from ``static_datas/makine_ekipman.xlsx`` by default. The
``Ekipman Listesi`` sheet is the owned-fleet list; the rented, returned, sold
and duplicate generator sheets are intentionally not imported.

Machines are matched by ``all_machines.identification_no``. Existing machines
are left untouched unless ``--overwrite`` is given. New machines are always
created with the ``Aktif`` status. Lookup tables (territory, type, mark, model,
company) are resolved by name; ``--create-lookups`` allows missing type,
subtype and mark values to be created automatically.
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl
from sqlalchemy import select

from database.setup import SessionLocal
from models.machines_model import (
    AllMachineModel,
    CarMarkModel,
    CarModelModel,
    CarStatusModel,
    CompanyModel,
    SubTypeTransportModel,
    TerritoryModel,
    TypeTransportModel,
)
from models.user_model import UserModel  # noqa: F401  (mapper dependency)

SERVER_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = SERVER_DIR / "static_datas" / "makine_ekipman.xlsx"

SHEET_NAME = "Ekipman Listesi"
FIRST_DATA_ROW = 4

# One-based column indexes in the "Ekipman Listesi" sheet.
COL_ID = 2
COL_SITE = 3
COL_TYPE = 4
COL_BRAND_MODEL = 6
COL_PLATE = 7
COL_VIN = 8
COL_TECHNICAL = 9
COL_YEAR = 10
COL_WEIGHT = 11
COL_DIMENSION = 12
COL_ENGINE_POWER = 13
COL_ENGINE_MARK_MODEL = 14
COL_ENGINE_IDENTITY = 15
COL_COMPANY = 16


# Cyrillic is visually transliterated so equipment codes such as С41К92 and
# ХЕ210WD match Latin lookups; names such as КАМАЗ and ГАЗ still normalize to
# KAMAZ and GAZ.
_CYRILLIC = str.maketrans(
    {
        "А": "A",
        "Б": "B",
        "В": "B",
        "Г": "G",
        "Д": "D",
        "Е": "E",
        "Ё": "E",
        "Ж": "ZH",
        "З": "Z",
        "И": "I",
        "Й": "I",
        "К": "K",
        "Л": "L",
        "М": "M",
        "Н": "H",
        "О": "O",
        "П": "P",
        "Р": "P",
        "С": "C",
        "Т": "T",
        "У": "Y",
        "Ф": "F",
        "Х": "X",
        "Ц": "TS",
        "Ч": "CH",
        "Ш": "SH",
        "Щ": "SHCH",
        "Ъ": "",
        "Ы": "Y",
        "Ь": "",
        "Э": "E",
        "Ю": "YU",
        "Я": "YA",
        "а": "a",
        "б": "b",
        "в": "b",
        "г": "g",
        "д": "d",
        "е": "e",
        "ё": "e",
        "ж": "zh",
        "з": "z",
        "и": "i",
        "й": "i",
        "к": "k",
        "л": "l",
        "м": "m",
        "н": "h",
        "о": "o",
        "п": "p",
        "р": "p",
        "с": "c",
        "т": "t",
        "у": "y",
        "ф": "f",
        "х": "x",
        "ц": "ts",
        "ч": "ch",
        "ш": "sh",
        "щ": "shch",
        "ъ": "",
        "ы": "y",
        "ь": "",
        "э": "e",
        "ю": "yu",
        "я": "ya",
    }
)
_TURKISH = str.maketrans("ıİşŞğĞçÇöÖüÜ", "iIsSgGcCoOuU")

# Workbook type labels whose name differs from the database lookup name.
TYPE_CANONICAL = {
    "binek arac": "Binek Arac",
    "otobus": "Otobus",
    "minibus": "Minibus",
    "kamyonet": "Kamyonet",
    "tir dorse": "Tir",
    "tir cekici": "Tir",
    "mobil vinc": "Mobil Vinc",
    "jenerator": "Jenerator",
    "isildakli jenerator": "Jenerator",
    "backhoe loader": "Backhoe-Loader",
    "forklift teleskopik": "Forklift",
    "forklift mini catalli mini forklift": "Forklift",
    "mini loader": "Mini Loader",
    "hi up": "Hi-Up",
    "excavator lastikli": "Excavator",
    "su tankeri": "Su Tankeri",
    "traktor": "Traktor",
    "eklemli bomlu platform manlift": "Manlift",
    "kompressor": "Kompressor",
    "vidanjor": "Vidanjör",
}

SUBTYPE_CANONICAL = {
    "tir dorse": "Dorse",
    "tir cekici": "Cekici",
    "forklift teleskopik": "Teleskopik",
    "forklift mini catalli mini forklift": "Çatallı Mini Forklift",
}

# Spelling differences between the workbook and the existing lookup values.
MODEL_ALIASES = {
    "ha32rtjpro": "ha32trjpro",
}

# Display names for marks that are not yet present in the lookup table.
MARK_DISPLAY_OVERRIDES = {
    "cobol": "Sobol",
    "sobol": "Sobol",
}


def _clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).replace("\u00a0", " ").strip()
    text = " ".join(text.split())
    if not text or text == "-":
        return None
    return text


def _normalize_id(value: Any) -> Optional[str]:
    text = _clean_text(value)
    return text.upper() if text else None


def _unique_key(value: Any) -> Optional[str]:
    """Case/space-insensitive key used for VIN and plate conflict checks."""
    text = _clean_text(value)
    if not text:
        return None
    return re.sub(r"\s+", "", text).upper()


def _normalize_text(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ""
    text = text.translate(_CYRILLIC).translate(_TURKISH)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.casefold()
    text = text.replace("mitsubushi", "mitsubishi")
    text = text.replace("hilight", "highlight")
    text = text.replace("eatchforce", "earthforce")
    text = text.replace("rm25", "pm25")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _parse_number(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).replace("\u00a0", "").replace(" ", "").strip()
    if not text or text.casefold() in {"-", "n/a"}:
        return None
    match = re.search(r"-?\d+(?:[.,]\d+)?", text)
    if match is None:
        return None
    try:
        return float(match.group().replace(",", "."))
    except ValueError:
        return None


def _parse_year(value: Any) -> Optional[datetime]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        year = value.year
    elif isinstance(value, (int, float)):
        year = int(value)
    else:
        match = re.search(r"\b(19|20)\d{2}\b", str(value))
        if match is None:
            return None
        year = int(match.group())

    if not 1900 <= year <= 2100:
        return None
    return datetime(year, 1, 1, tzinfo=timezone.utc)


def _cell(row: tuple[Any, ...], column: int) -> Any:
    index = column - 1
    return row[index] if index < len(row) else None


def read_machines(workbook) -> list[dict[str, Any]]:
    if SHEET_NAME not in workbook.sheetnames:
        raise SystemExit(f"Sheet not found: {SHEET_NAME}")

    ws = workbook[SHEET_NAME]
    records: list[dict[str, Any]] = []
    seen_ids: set[str] = set()

    for row_number, row in enumerate(
        ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True),
        start=FIRST_DATA_ROW,
    ):
        identification_no = _normalize_id(_cell(row, COL_ID))
        if not identification_no:
            continue

        if identification_no in seen_ids:
            print(
                f"  ! {identification_no} (row {row_number}): duplicate ID in "
                f"'{SHEET_NAME}', skipped"
            )
            continue
        seen_ids.add(identification_no)

        plate_no = _clean_text(_cell(row, COL_PLATE))
        vin_no = _clean_text(_cell(row, COL_VIN))
        if plate_no:
            plate_no = plate_no.upper()
        if vin_no:
            vin_no = vin_no.upper()

        year_raw = _cell(row, COL_YEAR)
        production_year = _parse_year(year_raw)
        if production_year is None and _clean_text(year_raw) is not None:
            print(
                f"  ! {identification_no}: unrecognised production year "
                f"{year_raw!r}, stored as empty"
            )

        records.append(
            {
                "identification_no": identification_no,
                "site": _clean_text(_cell(row, COL_SITE)),
                "type_label": _clean_text(_cell(row, COL_TYPE)),
                "brand_model": _clean_text(_cell(row, COL_BRAND_MODEL)),
                "plate_no": plate_no,
                "vin_no": vin_no,
                "technical_character": _clean_text(_cell(row, COL_TECHNICAL)),
                "production_year": production_year,
                "weight": _parse_number(_cell(row, COL_WEIGHT)),
                "dimension": _clean_text(_cell(row, COL_DIMENSION)),
                "engine_power": _clean_text(_cell(row, COL_ENGINE_POWER)),
                "engine_mark_model": _clean_text(
                    _cell(row, COL_ENGINE_MARK_MODEL)
                ),
                "engine_identity": _clean_text(_cell(row, COL_ENGINE_IDENTITY)),
                "company": _clean_text(_cell(row, COL_COMPANY)),
                "row": row_number,
            }
        )

    return records


def _find_model(
    remainder: str,
    models: dict[str, CarModelModel],
) -> Optional[CarModelModel]:
    if not remainder:
        return None

    model = models.get(remainder)
    if model is not None:
        return model

    alias = MODEL_ALIASES.get(remainder)
    if alias is not None:
        model = models.get(alias)
        if model is not None:
            return model

    # Unknown marks: still try to find the model at the end of the string.
    candidates = [
        (normalized, candidate)
        for normalized, candidate in models.items()
        if remainder == normalized or remainder.endswith(" " + normalized)
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda item: len(item[0]))[1]


async def _resolve_mark_and_model(
    session,
    brand_model: Optional[str],
    marks: list[tuple[str, CarMarkModel]],
    models: dict[str, CarModelModel],
    *,
    create_lookups: bool,
) -> tuple[Optional[CarMarkModel], Optional[CarModelModel]]:
    text = _normalize_text(brand_model)
    if not text:
        return None, None

    mark: Optional[CarMarkModel] = None
    remainder = text
    if text.startswith("gaz gazelle "):
        mark = next(
            (candidate for normalized, candidate in marks if normalized == "gazelle"),
            None,
        )
        if mark is not None:
            remainder = text[len("gaz gazelle ") :].strip()

    if mark is None:
        for normalized, candidate in marks:
            if text == normalized or text.startswith(normalized + " "):
                mark = candidate
                remainder = text[len(normalized) :].strip()
                break

    if mark is None and create_lookups:
        first, _, rest = text.partition(" ")
        display = MARK_DISPLAY_OVERRIDES.get(first)
        if display is None:
            original = (brand_model or "").split(None, 1)
            display = original[0] if original else first
        if display:
            mark = CarMarkModel(name=display)
            session.add(mark)
            await session.flush()
            marks.append((_normalize_text(display), mark))
            remainder = rest

    if mark is None:
        return None, _find_model(text, models)

    # ZIF uses "ПВ" in the workbook while the lookup stores "PB 10/10.1".
    if _normalize_text(mark.name) == "zif" and remainder.startswith(
        ("pv 10 1", "pb 10")
    ):
        model = models.get("pb 10 10 1")
        if model is not None:
            return mark, model

    # "GAZ GAZELLE NEXT ..." contains the mark twice.
    if remainder.startswith("gazelle "):
        remainder = remainder[len("gazelle ") :]

    return mark, _find_model(remainder, models)


async def _resolve_type(
    session,
    type_label: Optional[str],
    types: dict[str, TypeTransportModel],
    subtypes: dict[str, SubTypeTransportModel],
    *,
    create_lookups: bool,
) -> tuple[Optional[TypeTransportModel], Optional[SubTypeTransportModel]]:
    normalized = _normalize_text(type_label)
    if not normalized:
        return None, None

    canonical = TYPE_CANONICAL.get(normalized, type_label)
    type_obj = types.get(_normalize_text(canonical))

    subtype_obj: Optional[SubTypeTransportModel] = None
    subtype_name = SUBTYPE_CANONICAL.get(normalized)
    if subtype_name:
        subtype_obj = subtypes.get(_normalize_text(subtype_name))
        if subtype_obj is None and create_lookups:
            subtype_obj = SubTypeTransportModel(name=subtype_name)
            session.add(subtype_obj)
            await session.flush()
            subtypes[_normalize_text(subtype_name)] = subtype_obj

    if type_obj is None and create_lookups and canonical:
        type_obj = TypeTransportModel(name=canonical)
        session.add(type_obj)
        await session.flush()
        types[_normalize_text(canonical)] = type_obj

    return type_obj, subtype_obj


def _add_gap(
    gaps: dict[str, list[str]],
    field: str,
    machine_id: str,
    value: Any,
) -> None:
    if value:
        gaps.setdefault(field, []).append(f"{machine_id} ({value})")


def _print_summary(
    records: list[dict[str, Any]],
    *,
    created: int,
    updated: int,
    skipped: int,
    conflicts: list[str],
    gaps: dict[str, list[str]],
) -> None:
    print()
    print(f"Machines found in Excel : {len(records)}")
    print(f"  new                   : {created}")
    print(f"  updated               : {updated}")
    print(f"  skipped (existing)    : {skipped}")
    print(f"  skipped (conflict)    : {len(conflicts)}")
    if conflicts:
        print("  " + "; ".join(conflicts))

    if gaps:
        print("\nLookup values not matched:")
        for field, items in gaps.items():
            shown = ", ".join(items[:25])
            if len(items) > 25:
                shown += f" ... (+{len(items) - 25} more)"
            print(f"  {field}: {len(items)}")
            print(f"    {shown}")


async def seed_database(
    records: list[dict[str, Any]],
    *,
    dry_run: bool,
    overwrite: bool,
    create_lookups: bool,
) -> None:
    async with SessionLocal() as session:
        territory_rows = (await session.execute(select(TerritoryModel))).scalars()
        type_rows = (await session.execute(select(TypeTransportModel))).scalars()
        subtype_rows = (
            await session.execute(select(SubTypeTransportModel))
        ).scalars()
        mark_rows = (await session.execute(select(CarMarkModel))).scalars()
        model_rows = (await session.execute(select(CarModelModel))).scalars()
        company_rows = (await session.execute(select(CompanyModel))).scalars()
        status_rows = (await session.execute(select(CarStatusModel))).scalars()

        territories = {
            _normalize_text(row.name): row
            for row in territory_rows
            if _normalize_text(row.name)
        }
        types = {
            _normalize_text(row.name): row
            for row in type_rows
            if _normalize_text(row.name)
        }
        subtypes = {
            _normalize_text(row.name): row
            for row in subtype_rows
            if _normalize_text(row.name)
        }
        models = {
            _normalize_text(row.name): row
            for row in model_rows
            if _normalize_text(row.name)
        }
        companies = {
            _normalize_text(row.name): row
            for row in company_rows
            if _normalize_text(row.name)
        }
        marks = sorted(
            (
                (_normalize_text(row.name), row)
                for row in mark_rows
                if _normalize_text(row.name)
            ),
            key=lambda item: len(item[0]),
            reverse=True,
        )

        active_status = next(
            (
                row
                for row in status_rows
                if _normalize_text(row.name) == "aktif"
            ),
            None,
        )
        if active_status is None:
            raise SystemExit(
                "Active status 'Aktif' was not found in car_statuses; "
                "run the status seed first."
            )

        existing_rows = (
            await session.execute(select(AllMachineModel))
        ).scalars().all()
        existing_by_id = {
            normalized: machine
            for machine in existing_rows
            if (normalized := _normalize_id(machine.identification_no))
        }
        vin_owners = {
            key: machine
            for machine in existing_rows
            if (key := _unique_key(machine.vin_no))
        }
        plate_owners = {
            key: machine
            for machine in existing_rows
            if (key := _unique_key(machine.plate_no))
        }

        created = 0
        updated = 0
        skipped = 0
        conflicts: list[str] = []
        gaps: dict[str, list[str]] = {}

        for record in records:
            machine_id = record["identification_no"]
            existing = existing_by_id.get(machine_id)

            if existing is not None and not overwrite:
                skipped += 1
                continue

            vin_key = _unique_key(record["vin_no"])
            plate_key = _unique_key(record["plate_no"])
            vin_owner = vin_owners.get(vin_key) if vin_key else None
            plate_owner = plate_owners.get(plate_key) if plate_key else None

            if vin_owner is not None and vin_owner is not existing:
                conflicts.append(
                    f"{machine_id}: VIN {record['vin_no']} already belongs to "
                    f"{vin_owner.identification_no}"
                )
                continue
            if plate_owner is not None and plate_owner is not existing:
                conflicts.append(
                    f"{machine_id}: plate {record['plate_no']} already belongs "
                    f"to {plate_owner.identification_no}"
                )
                continue

            territory = territories.get(_normalize_text(record["site"]))
            if territory is None:
                _add_gap(gaps, "territory", machine_id, record["site"])

            type_obj, subtype_obj = await _resolve_type(
                session,
                record["type_label"],
                types,
                subtypes,
                create_lookups=create_lookups,
            )
            if type_obj is None:
                _add_gap(gaps, "type", machine_id, record["type_label"])
            if subtype_obj is None and record["type_label"] in {
                "Tır Dorse",
                "Tır (Çekici)",
                "Forklift (Teleskopik)",
                "Forklift Mini (Çatallı Mini Forklift)",
            }:
                _add_gap(gaps, "subtype", machine_id, record["type_label"])

            mark, model = await _resolve_mark_and_model(
                session,
                record["brand_model"],
                marks,
                models,
                create_lookups=create_lookups,
            )
            if mark is None:
                _add_gap(gaps, "mark", machine_id, record["brand_model"])
            if model is None:
                _add_gap(gaps, "model", machine_id, record["brand_model"])

            company = companies.get(_normalize_text(record["company"]))
            if company is None and record["company"]:
                _add_gap(gaps, "company", machine_id, record["company"])

            if existing is not None:
                old_vin_key = _unique_key(existing.vin_no)
                if old_vin_key and vin_owners.get(old_vin_key) is existing:
                    vin_owners.pop(old_vin_key, None)
                old_plate_key = _unique_key(existing.plate_no)
                if old_plate_key and plate_owners.get(old_plate_key) is existing:
                    plate_owners.pop(old_plate_key, None)
                machine = existing
                updated += 1
            else:
                machine = AllMachineModel(
                    identification_no=machine_id,
                    status_id=active_status.id,
                )
                session.add(machine)
                created += 1

            machine.identification_no = machine_id
            machine.vin_no = record["vin_no"]
            machine.plate_no = record["plate_no"]
            machine.technical_character = record["technical_character"]
            machine.production_year = record["production_year"]
            machine.weight = record["weight"]
            machine.dimension = record["dimension"]
            machine.engine_power = record["engine_power"]
            machine.engine_mark_model = record["engine_mark_model"]
            machine.engine_identity = record["engine_identity"]
            machine.territory_id = territory.id if territory else None
            machine.type_id = type_obj.id if type_obj else None
            machine.subtype_id = subtype_obj.id if subtype_obj else None
            machine.car_mark_id = mark.id if mark else None
            machine.car_model_id = model.id if model else None
            machine.company_id = company.id if company else None

            if existing is None:
                machine.status_id = active_status.id

            if vin_key:
                vin_owners[vin_key] = machine
            if plate_key:
                plate_owners[plate_key] = machine

        _print_summary(
            records,
            created=created,
            updated=updated,
            skipped=skipped,
            conflicts=conflicts,
            gaps=gaps,
        )

        if dry_run:
            await session.rollback()
            print("\nDry run: no changes were saved.")
            return

        await session.commit()
        print("\nMachines saved.")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Seed the machine/equipment fleet from the Excel workbook."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help=f"Path to the .xlsx file (default: {DEFAULT_WORKBOOK})",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Read the workbook and report what would change without saving.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Update machines that already exist (status is left unchanged).",
    )
    parser.add_argument(
        "--create-lookups",
        action="store_true",
        help="Create missing type, subtype and mark lookup values.",
    )
    args = parser.parse_args()

    workbook_path: Path = args.workbook
    if not workbook_path.is_absolute():
        workbook_path = SERVER_DIR / workbook_path
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    print(f"Reading workbook: {workbook_path}")
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    try:
        records = read_machines(workbook)
    finally:
        workbook.close()

    try:
        asyncio.run(
            seed_database(
                records,
                dry_run=args.dry_run,
                overwrite=args.overwrite,
                create_lookups=args.create_lookups,
            )
        )
    except KeyboardInterrupt:
        sys.exit(130)


if __name__ == "__main__":
    main()
